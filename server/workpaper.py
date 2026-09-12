"""調書グリッドのサーバ側（IronCalc Python）。FastAPI の APIRouter として他アプリに include できる。

  from workpaper import router
  app.include_router(router, prefix="/api")

役割: xlsx→IronCalcバイナリ変換 / ブラウザからの差分適用 / xlsx保存（変更履歴シート・結合の反映）/ 編集ログの保管 / AI指示の窓口
状態はモジュール内のメモリ（試作）。本番は SESSIONS を永続化し、LOGS は DB に、AI は既存ロジックへ差し替える。
"""
import datetime, hashlib, json, os, re, tempfile, uuid
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
import ironcalc as ic
import openpyxl

from . import xlsx_pkg

router = APIRouter()

TMPDIR = tempfile.gettempdir()   # Windows でも動くよう /tmp 直書きを避ける
LOG_DIR = os.path.join(TMPDIR, "workpaper_logs")  # 試作: セッションごとの追記専用ファイル（本番はDB）
os.makedirs(LOG_DIR, exist_ok=True)
LOG_SHEET = "変更履歴"  # 保存時に xlsx へ同梱するシート名

SESSIONS: dict[str, ic.UserModel] = {}   # 試作用の常駐モデル（本番はDB/ファイルに永続化）
ORIGINALS: dict[str, bytes] = {}         # 開いた元の xlsx（保存時の土台。図形・メモ等はここから残す）
LOGS: dict[str, list[dict]] = {}         # セッションごとの編集ログ（追記のみ）


# ---- 開く ------------------------------------------------------------------
@router.post("/open")
async def open_workbook(file: UploadFile = File(...)):
    data = await file.read()
    path = os.path.join(TMPDIR, f"{uuid.uuid4()}.xlsx")
    with open(path, "wb") as f:
        f.write(xlsx_pkg.strip_for_engine(data))  # IronCalc はメモ付きブックを開けないので、メモを除いたコピーを渡す
    model = ic.create_user_model_from_xlsx(path, "en", "UTC")
    model.evaluate()
    sid = uuid.uuid4().hex
    SESSIONS[sid] = model
    ORIGINALS[sid] = data
    LOGS[sid] = []
    # 結合セルは wasm API から取れないので、表示用に openpyxl で読んで渡す（[sheet, r1, c1, r2, c2]、1始まり）
    merges = []
    wb = openpyxl.load_workbook(path)
    for si, ws in enumerate(wb.worksheets):
        for rg in ws.merged_cells.ranges:
            merges.append([si, rg.min_row, rg.min_col, rg.max_row, rg.max_col])
    return Response(content=bytes(model.to_bytes()), media_type="application/octet-stream",
                    headers={"X-Session-Id": sid, "X-Merges": json.dumps(merges),
                             "Access-Control-Expose-Headers": "X-Session-Id, X-Merges"})


@router.get("/assets/{sid}")
async def get_assets(sid: str):
    """元ファイルにあって IronCalc が持たないもの（メモ・画像・グラフ・図形・リンク・入力規則）を画面表示用に返す"""
    if sid not in ORIGINALS:
        raise HTTPException(404, "session not found")
    data = ORIGINALS[sid]
    return {"notes": xlsx_pkg.read_notes(data), **xlsx_pkg.read_assets(data)}


# ---- 編集ログ ------------------------------------------------------------
class LogCell(BaseModel):
    addr: str          # 例: 手続!C8
    before: str = ""
    after: str = ""
    note: str = ""     # AI更新なら根拠になったコメント


class LogEntry(BaseModel):
    seq: int           # ブラウザ側の操作番号（undo/redo の ref に使う）
    ts: str            # ブラウザ側の時刻
    actor: str         # user / ai
    action: str        # edit / paste / clear / insert_rows / ... / undo / redo / save（web/src/workpaper-grid/log.ts と同じ）
    sheet: str
    range: str
    cells: list[LogCell] = []
    note: str = ""
    ref: int | None = None


class LogRequest(BaseModel):
    entries: list[LogEntry]


@router.post("/log/{sid}")
async def append_log(sid: str, req: LogRequest):
    """編集ログを追記する（削除・更新はしない）。
    受信時刻と、直前の行のハッシュを含めたハッシュを各行に付け、後から差し替えられないようにする。"""
    if sid not in SESSIONS:
        raise HTTPException(404, "session not found")
    log = LOGS.setdefault(sid, [])
    with open(os.path.join(LOG_DIR, f"{sid}.jsonl"), "a", encoding="utf-8") as f:
        for e in req.entries:
            rec = e.model_dump()
            rec["server_ts"] = datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")
            prev = log[-1]["hash"] if log else ""
            rec["hash"] = hashlib.sha256((prev + json.dumps(rec, ensure_ascii=False, sort_keys=True)).encode()).hexdigest()[:16]
            log.append(rec)
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return {"count": len(log)}


@router.get("/log/{sid}")
async def get_log(sid: str):
    if sid not in SESSIONS:
        raise HTTPException(404, "session not found")
    return {"entries": LOGS.get(sid, [])}


def _text(v: str) -> str:
    """数式として解釈されないよう、= で始まる値は文字列扱いにする"""
    return "'" + v if v.startswith("=") else v


def write_log_sheet(model: ic.UserModel, log: list[dict]) -> None:
    """変更履歴シートを作り直して書く。調書本体と履歴が同じ xlsx に残る。"""
    names = [w["name"] for w in model.get_worksheets_properties()]
    if LOG_SHEET in names:
        idx = names.index(LOG_SHEET)
        model.range_clear_all(idx, 1, 1, 5000, 10)
    else:
        model.new_sheet()
        idx = len(names)
        model.rename_sheet(idx, LOG_SHEET)
    headers = ["#", "日時(UTC)", "操作者", "操作", "シート", "セル/範囲", "変更前", "変更後", "備考", "ハッシュ"]
    widths = [40, 150, 60, 100, 90, 110, 220, 220, 220, 130]
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        model.set_user_input(idx, 1, c, h)
        model.set_columns_width(idx, c, c, w)
    r = 2
    for e in log:
        for cell in e["cells"] or [{}]:   # 1セル1行に展開
            row = [str(e["seq"]), e["server_ts"], "AI" if e["actor"] == "ai" else "手動", e["action"], e["sheet"],
                   cell.get("addr") or e["range"], _text(cell.get("before", "")), _text(cell.get("after", "")),
                   cell.get("note") or e["note"], e["hash"]]
            for c, v in enumerate(row, 1):
                if v:
                    model.set_user_input(idx, r, c, v)
            r += 1


def apply_merges(path: str, merges: list[list[int]]) -> None:
    """ブラウザで変更した結合状態を xlsx に反映する（IronCalc の wasm に結合 API が無いため openpyxl で後処理）。
    merges は [sheet, r1, c1, r2, c2] の一覧で、変更履歴シートより前のシートが対象。"""
    wb = openpyxl.load_workbook(path)
    by_sheet: dict[int, list[list[int]]] = {}
    for si, r1, c1, r2, c2 in merges:
        by_sheet.setdefault(si, []).append([r1, c1, r2, c2])
    for si, ws in enumerate(wb.worksheets):
        if ws.title == LOG_SHEET:
            continue
        for rg in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rg))
        for r1, c1, r2, c2 in by_sheet.get(si, []):
            if r2 > r1 or c2 > c1:
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    wb.save(path)


# ---- 保存 ------------------------------------------------------------------
@router.post("/sync/{sid}")
async def sync(sid: str, file: UploadFile = File(...), merges: str = Form(""), orig_names: str = Form(""), struct_ops: str = Form("")):
    """ブラウザ側 flushSendQueue() の差分をサーバ側モデルに適用し、xlsx を返す。
    orig_names（各シートが元ファイルのどのシートか）が来れば、元ファイルを土台に IronCalc の出力からセル部分だけを移植して
    図形・メモ・フィルタ・入力規則・リンクを残す（xlsx_pkg.graft）。来なければ IronCalc の出力をそのまま返す（従来動作）。"""
    model = SESSIONS.get(sid)
    if model is None:
        raise HTTPException(404, "session not found")
    diffs = await file.read()
    if diffs:
        model.apply_external_diffs(diffs)
    write_log_sheet(model, LOGS.get(sid, []))
    out = os.path.join(TMPDIR, f"{sid}-{uuid.uuid4().hex}.xlsx")  # save_to_xlsx は既存ファイルを上書きしないため毎回別名
    model.save_to_xlsx(out)
    merge_list = json.loads(merges) if merges else []
    if orig_names and sid in ORIGINALS:
        with open(out, "rb") as f:
            engine = f.read()
        data = xlsx_pkg.graft(ORIGINALS[sid], engine, json.loads(orig_names), merge_list, json.loads(struct_ops) if struct_ops else [])
    else:
        if merge_list:
            apply_merges(out, merge_list)
        with open(out, "rb") as f:
            data = f.read()
    os.remove(out)
    return Response(content=data,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---- AI（ダミー） ----------------------------------------------------------
class AiItem(BaseModel):
    """コメントが付いたセル1件。current は現在の内容、comment はユーザーの修正指示"""
    sheet: int
    row: int
    col: int
    current: str
    comment: str


class AiRequest(BaseModel):
    items: list[AiItem]


@router.post("/ai")
async def ai_edit(req: AiRequest):
    """コメント付きセルをまとめて受け取り、セルごとの操作命令 ops を返す（部分再生成）。
    本番では既存のAIロジックへ委譲する。ここでは形式を示すためのダミー:
    数値指示なら数値、それ以外は文章として提案する。"""
    ops = []
    for it in req.items:
        m = re.search(r"-?\d+(\.\d+)?", it.comment)
        if m and any(k in it.comment for k in ("件", "に", "を", "set")):
            value = m.group(0)
        else:
            value = f"{it.current}\n[AI案] {it.comment}".strip()
        ops.append({"sheet": it.sheet, "row": it.row, "col": it.col, "input": value})
    return {"ops": ops,
            "note": "ダミー応答。実運用では既存AIロジックの出力を ops 形式に整形して返す。"}
