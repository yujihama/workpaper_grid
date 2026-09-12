from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb = Workbook()
ws = wb.active; ws.title = "手続"
thin = Side(style="thin"); b = Border(left=thin,right=thin,top=thin,bottom=thin)
ws.merge_cells("A1:E1"); ws["A1"] = "内部監査調書：購買プロセス（J-SOX）"; ws["A1"].font = Font(bold=True, size=14)
ws["A3"]="調書番号"; ws["B3"]="PUR-001"; ws["D3"]="作成者"; ws["E3"]=""
ws["A4"]="対象期間"; ws["B4"]="FY2026 Q1"; ws["D4"]="作成日"; ws["E4"]=""
hdr = ["No","手続","実施結果","判断","証跡"]
for i,h in enumerate(hdr,1):
    c = ws.cell(row=6,column=i,value=h); c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDE4EE"); c.border=b
rows = [(1,"発注書の承認権限者による承認を確認する","","",""),
        (2,"検収記録と請求書の突合を実施する","","",""),
        (3,"サンプル25件の三点照合を行う","","","")]
for r,row in enumerate(rows,7):
    for c,v in enumerate(row,1):
        cell=ws.cell(row=r,column=c,value=v); cell.border=b; cell.alignment=Alignment(wrap_text=True,vertical="top")
ws["A11"]="不備件数（サンプリングより）"; ws["B11"]="=サンプリング!B6"
ws["A12"]="総合判断"; ws["B12"]='=IF(B11=0,"不備なし",IF(B11<=2,"要検討","不備あり"))'
ws.column_dimensions["B"].width=40; ws.column_dimensions["C"].width=40
ws2 = wb.create_sheet("サンプリング")
ws2["A1"]="母集団件数"; ws2["B1"]=1200
ws2["A2"]="サンプル数"; ws2["B2"]=25
ws2["A3"]="抽出率"; ws2["B3"]="=B2/B1"; ws2["B3"].number_format="0.00%"
ws2["A5"]="No"; ws2["B5"]="結果(1=例外)"
ws2["A6"]="例外件数"; ws2["B6"]="=SUM(B8:B32)"
for i in range(25):
    ws2.cell(row=8+i,column=1,value=i+1); ws2.cell(row=8+i,column=2,value=0)
ws2["B9"]=1
wb.save("proto/sample_workpaper.xlsx"); print("ok")
